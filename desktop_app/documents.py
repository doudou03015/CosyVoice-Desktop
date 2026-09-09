"""Document import and isolated, read-only PowerPoint rendering."""
from __future__ import annotations

import hashlib
import csv
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import threading
import time
from uuid import uuid4

from .paths import app_root, outside_sync, session_dir


def sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_text(path):
    path = Path(path)
    if path.suffix.lower() == ".docx":
        from docx import Document
        # Paragraphs and table cells remain in source document order.
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(path)
        blocks = []
        for node in document.element.body:
            if node.tag.endswith("}p"):
                blocks.append(Paragraph(node, document).text)
            elif node.tag.endswith("}tbl"):
                blocks.extend("\t".join(cell.text for cell in row.cells)
                              for row in Table(node, document).rows)
        return "\n".join(blocks).strip()
    if path.suffix.lower() != ".txt":
        raise ValueError("请选择 TXT 或 DOCX 文件。")
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "gb18030"):
        try:
            return data.decode(encoding).replace("\r\n", "\n").replace("\r", "\n").strip()
        except UnicodeError:
            continue
    raise ValueError("无法读取讲稿编码，请另存为 UTF-8 TXT。")


def parse_page_script(text, page_count):
    if page_count < 1:
        raise ValueError("PPT 中没有页面。")
    result, current, buffer = {}, None, []
    def flush():
        if current is not None:
            result[current] = "\n".join(buffer).strip()
    for line in text.replace("\r\n", "\n").split("\n"):
        match = re.fullmatch(r"\s*第\s*(\d+)\s*页\s*[:：]?\s*", line)
        if match:
            flush()
            current = int(match.group(1))
            if not 1 <= current <= page_count:
                raise ValueError(f"第 {current} 页超出 PPT 页码范围（1–{page_count}）。")
            if current in result:
                raise ValueError(f"第 {current} 页重复出现，请合并该页讲稿。")
            buffer = []
        elif current is None:
            if line.strip():
                raise ValueError("讲稿须以独立一行的“第1页”开始。")
        else:
            buffer.append(line)
    flush()
    if not result:
        raise ValueError("未找到页码标题，请使用“第1页”“第2页”模板。")
    return result


def read_pptx(path):
    from pptx import Presentation
    path = Path(path).resolve()
    if path.suffix.lower() != ".pptx":
        raise ValueError("首版支持 PPTX，请先将旧格式另存为 .pptx。")
    before = sha256(path)
    presentation = Presentation(str(path))
    slides = []
    for number, slide in enumerate(presentation.slides, 1):
        text = ""
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                text = frame.text.strip()
        hidden = slide._element.get("show", "1").lower() in {"0", "false", "off"}
        title = slide.shapes.title.text if slide.shapes.title is not None else f"第 {number} 页"
        slides.append(dict(number=number, title=title, text=text, hidden=hidden,
                           included=not hidden, blank_seconds=3.0, voice_id="", speed=None,
                           status="pending", audio="", image=""))
    if before != sha256(path):
        raise RuntimeError("读取过程中 PPT 文件发生变化，请重新导入。")
    return dict(source=str(path), source_sha256=before, width=presentation.slide_width,
                height=presentation.slide_height, slides=slides)


def powerpoint_available():
    if os.name != "nt":
        return False
    try:
        import winreg
        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"PowerPoint.Application\CLSID"):
            return True
    except OSError:
        return False


def _render_event(destination, event):
    """Windowed executables may have no Python stdout/stderr at all."""
    with (Path(destination) / "renderer-events.jsonl").open("a", encoding="utf-8", newline="\n") as stream:
        stream.write(json.dumps(event, ensure_ascii=False) + "\n")
        stream.flush()


def render_cli(arguments):
    """Executed only by a dedicated child; never attaches to the user's PowerPoint."""
    source, destination = map(Path, arguments[:2])
    application = presentation = None
    owns_application = False
    try:
        import pythoncom
        import win32com.client
        def powerpoint_pids():
            output = subprocess.run(["tasklist", "/FI", "IMAGENAME eq POWERPNT.EXE", "/FO", "CSV", "/NH"],
                                    capture_output=True, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            return {int(row[1]) for row in csv.reader(io.StringIO(output.stdout.decode("utf-8", "replace")))
                    if len(row) >= 2 and row[0].lower() == "powerpnt.exe" and row[1].isdigit()}
        before_pids = powerpoint_pids()
        pythoncom.CoInitialize()
        application = win32com.client.DispatchEx("PowerPoint.Application")
        new_pids = powerpoint_pids() - before_pids
        # Application.HWND is unavailable in some Office versions. Only record an
        # unambiguously new PowerPoint pid; never kill by executable name.
        if len(new_pids) == 1:
            owns_application = True
            _render_event(destination, {"event": "powerpoint_pid", "pid": new_pids.pop()})
        else:
            raise RuntimeError("未能确认独立的 PowerPoint 实例，请保存并关闭 PowerPoint 后重试。")
        application.DisplayAlerts = 1
        application.AutomationSecurity = 3  # Disable macros for automation documents.
        presentation = application.Presentations.Open(str(source.resolve()), ReadOnly=-1,
                                                     Untitled=0, WithWindow=0)
        width, height = presentation.PageSetup.SlideWidth, presentation.PageSetup.SlideHeight
        export_width = round(1920 if width / height >= 1920 / 1080 else 1080 * width / height)
        export_height = round(export_width * height / width)
        for index in range(1, presentation.Slides.Count + 1):
            output = destination / f"slide_{index:04d}.png"
            presentation.Slides(index).Export(str(output), "PNG", export_width, export_height)
            _render_event(destination, {"event": "progress", "stage": "render", "completed": index,
                                       "total": presentation.Slides.Count,
                                       "message": f"已渲染第 {index} 页"})
        return 0
    except Exception as error:
        _render_event(destination, {"event": "error", "message": str(error)})
        return 1
    finally:
        if presentation is not None:
            try:
                presentation.Close()
            except Exception:
                pass
        if application is not None and owns_application:
            try:
                application.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except (NameError, AttributeError):
            pass


def render_pptx(path, output_dir, progress=None, cancel=None):
    if not powerpoint_available():
        raise RuntimeError("未检测到 Microsoft PowerPoint。请安装后使用 PPT 视频功能。")
    path = Path(path).resolve()
    before = sha256(path)
    temporary = outside_sync(session_dir() / ("render-" + uuid4().hex))
    temporary.mkdir()
    # PowerPoint opens a working copy, so even Office lock/recovery files stay in Temp.
    working = temporary / "source.pptx"
    shutil.copyfile(path, working)
    args = [sys.executable]
    if not getattr(sys, "frozen", False):
        args += ["-B", "-m", "desktop_app.documents"]
    args += ["--powerpoint-render", str(working), str(temporary)]
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    child = subprocess.Popen(args, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, text=True,
                             encoding="utf-8", errors="replace", creationflags=flags,
                             cwd=str(app_root()))
    errors = []
    def pump(stream):
        for line in stream:
            errors.append(line)
    threads = [threading.Thread(target=pump, args=(child.stderr,), daemon=True)]
    for thread in threads:
        thread.start()
    office_pid = None
    started = time.monotonic()
    event_path = temporary / "renderer-events.jsonl"
    offset = 0
    pending = b""
    try:
        while True:
            if cancel and cancel():
                raise InterruptedError("已取消页面渲染。")
            if time.monotonic() - started > 1800:
                raise RuntimeError("PowerPoint 渲染超时，请检查文档或 Office 弹窗。")
            if event_path.is_file():
                with event_path.open("rb") as stream:
                    stream.seek(offset)
                    block = stream.read()
                    offset += len(block)
                    pending += block
                while b"\n" in pending:
                    line, _, pending = pending.partition(b"\n")
                    try:
                        event = json.loads(line.decode("utf-8"))
                    except (ValueError, UnicodeError):
                        continue
                    if event.get("event") == "powerpoint_pid":
                        office_pid = event["pid"]
                    elif event.get("event") == "error":
                        errors.append(event["message"])
                    elif progress:
                        progress(event)
            if child.poll() is not None:
                # Read once more after exit; the last event may have been flushed
                # between this iteration's file read and process termination.
                if event_path.is_file() and event_path.stat().st_size > offset:
                    continue
                break
            time.sleep(0.05)
        for thread in threads:
            thread.join(timeout=1)
        if child.wait() != 0:
            raise RuntimeError("PowerPoint 页面渲染失败：" + "".join(errors)[-1500:])
        if before != sha256(path):
            raise RuntimeError("渲染期间原始 PPT 被外部修改，请重新导入。")
        destination = Path(output_dir).resolve()
        destination.mkdir(parents=True, exist_ok=True)
        result = {}
        for file in sorted(temporary.glob("slide_*.png")):
            target = destination / file.name
            shutil.copyfile(file, target)
            result[int(file.stem.split("_")[-1])] = str(target)
        if not result:
            raise RuntimeError("PowerPoint 未返回页面图像。")
        return result
    finally:
        if child.poll() is None:
            child.kill()
            child.wait(timeout=10)
            # DispatchEx pid belongs exclusively to this renderer, never a name-based kill.
            if office_pid and os.name == "nt":
                subprocess.run(["taskkill", "/PID", str(office_pid), "/F"],
                               capture_output=True, creationflags=flags)
        for thread in threads:
            thread.join(timeout=1)
        child.stderr.close()
        shutil.rmtree(temporary, ignore_errors=True)


if __name__ == "__main__":
    if sys.argv[1:2] == ["--powerpoint-render"]:
        raise SystemExit(render_cli(sys.argv[2:]))
