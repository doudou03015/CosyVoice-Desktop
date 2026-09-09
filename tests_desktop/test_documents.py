import os
from pathlib import Path
import tempfile
import unittest
import json
import sys
from types import ModuleType, SimpleNamespace
from unittest.mock import Mock, patch

from desktop_app.documents import parse_page_script, read_text, read_pptx, render_cli, sha256
from desktop_app.paths import session_dir


def create_presentation(path, count=20, portrait=False):
    from pptx import Presentation
    from pptx.util import Inches
    presentation = Presentation()
    presentation.slide_width = Inches(7.5 if portrait else 13.3333)
    presentation.slide_height = Inches(10 if portrait else 7.5)
    for number in range(1, count + 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"第 {number} 页：中文课件"
        slide.placeholders[1].text = "日期 2026年9月9日，数值 123.45，English text。"
        if number % 5:
            slide.notes_slide.notes_text_frame.text = ("本页讲解科学方法和知识分享。" * (100 if number == 7 else 1))
        if number == 4:
            slide._element.set("show", "0")
        if number == 8:
            table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(5), Inches(1)).table
            table.cell(0, 0).text = "项目"
            table.cell(0, 1).text = "数值"
            table.cell(1, 0).text = "测试"
            table.cell(1, 1).text = "42"
    presentation.save(path)
    return path


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="documents-test-", dir=session_dir())
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_twenty_page_import_does_not_change_source(self):
        path = create_presentation(self.root / "含 空格课件.pptx")
        before = sha256(path)
        result = read_pptx(path)
        self.assertEqual(len(result["slides"]), 20)
        self.assertFalse(result["slides"][3]["included"])
        self.assertTrue(result["slides"][3]["hidden"])
        self.assertEqual(result["slides"][4]["text"], "")
        self.assertGreater(len(result["slides"][6]["text"]), 500)
        self.assertEqual(sha256(path), before)
        portrait = read_pptx(create_presentation(self.root / "竖版.pptx", 1, portrait=True))
        self.assertLess(portrait["width"], portrait["height"])

    def test_page_script_mapping_validation_and_blank_override(self):
        self.assertEqual(parse_page_script("第1页\n第一段\n第二段\n\n第 3 页：\n", 20),
                         {1: "第一段\n第二段", 3: ""})
        for text in ("第1页\n一\n第1页\n二", "第21页\n越界", "第0页\n错误", "无页码正文", ""):
            with self.subTest(text=text), self.assertRaises(ValueError):
                parse_page_script(text, 20)

    def test_txt_docx_and_tables(self):
        from docx import Document
        text = "普通文字\n第二段。"
        for encoding in ("utf-8-sig", "utf-16", "gb18030"):
            path = self.root / "讲稿.txt"
            path.write_text(text, encoding=encoding)
            self.assertEqual(read_text(path), text)
        doc = Document()
        doc.add_paragraph("第1页")
        doc.add_table(1, 2).cell(0, 0).text = "表格讲稿"
        doc.add_paragraph("末段")
        path = self.root / "讲稿.docx"
        doc.save(path)
        value = read_text(path)
        self.assertLess(value.index("第1页"), value.index("表格讲稿"))
        self.assertLess(value.index("表格讲稿"), value.index("末段"))

    def test_windowed_render_helper_reports_without_stdout_or_stderr(self):
        def export(path, format, width, height):
            Path(path).write_bytes(b"completed-image")
        slides = Mock(side_effect=lambda index: SimpleNamespace(Export=export))
        slides.Count = 2
        presentation = SimpleNamespace(Slides=slides,
            PageSetup=SimpleNamespace(SlideWidth=960, SlideHeight=540), Close=Mock())
        application = SimpleNamespace(Presentations=SimpleNamespace(Open=Mock(return_value=presentation)), Quit=Mock())
        pythoncom = SimpleNamespace(CoInitialize=Mock(), CoUninitialize=Mock())
        win32com = ModuleType("win32com")
        win32com.client = SimpleNamespace(DispatchEx=Mock(return_value=application))
        modules = {"pythoncom": pythoncom, "win32com": win32com, "win32com.client": win32com.client}
        with patch.dict(sys.modules, modules), patch("sys.stdout", None), patch("sys.stderr", None), \
             patch("desktop_app.documents.subprocess.run", side_effect=[
                 SimpleNamespace(stdout=b""), SimpleNamespace(stdout=b'"POWERPNT.EXE","12345"\n')]):
            result = render_cli([str(self.root / "source.pptx"), str(self.root)])
        self.assertEqual(result, 0)
        events = [json.loads(line) for line in (self.root / "renderer-events.jsonl").read_text(encoding="utf-8").splitlines()]
        self.assertEqual(events[0], {"event": "powerpoint_pid", "pid": 12345})
        self.assertEqual([event["completed"] for event in events[1:]], [1, 2])
        application.Quit.assert_called_once()
        presentation.Close.assert_called_once()


if __name__ == "__main__":
    unittest.main()
