"""Desktop wiring tests: real documents/projects/voices, simulated GPU completion."""
import os
from pathlib import Path
import shutil
import sys
import time

if os.name == "nt":
    import ctypes
    ctypes.windll.kernel32.SetErrorMode(0x0001 | 0x0002)

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from PySide6.QtCore import QCoreApplication, QEvent, Qt
from PySide6.QtGui import QFont
from PySide6.QtWidgets import QApplication
from pptx import Presentation

from desktop_app import projects
from desktop_app.ui import MainWindow


@pytest.fixture(scope="module")
def app():
    application = QApplication.instance() or QApplication([])
    application.setFont(QFont("Microsoft YaHei UI", 10))
    yield application


@pytest.fixture
def window(app, tmp_path, monkeypatch):
    monkeypatch.setenv("COSYVOICE_DESKTOP_DATA", str(tmp_path / "data"))
    monkeypatch.setattr(MainWindow, "detect_hardware", lambda self: None)
    result = MainWindow()
    result._last_gpu_probe = time.monotonic()
    errors = []
    monkeypatch.setattr(result, "show_error", errors.append)
    monkeypatch.setattr(result, "_ensure_inference_ready", lambda callback: True)
    result.test_errors = errors
    yield result
    result._dirty = False
    result.worker.active = None
    result.close()
    result.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
    app.processEvents()


def make_project(window, tmp_path):
    deck = Presentation()
    for number in range(3):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"测试页 {number + 1}"
        slide.notes_slide.notes_text_frame.text = f"第 {number + 1} 页的原始讲稿。"
    source = tmp_path / "中文 含空格.pptx"
    deck.save(source)
    folder = tmp_path / "工程 中文"
    project = projects.create_project(source, folder, voice=window.voice_items[0])
    window._set_project(project, folder)
    return project, folder


def test_four_sections_and_shared_builtin_voices(window):
    assert window.stack.count() == 4
    assert len(window.voice_items) == 6
    assert window.text_voice.current()["audio"] == window.ppt_voice.current()["audio"]
    assert window.text_seed.value() == 0
    assert window.text_speed.value() == 1
    assert window.ppt_tail.value() == .5
    # Importing the GUI must never initialize the inference runtime.
    assert "torch" not in sys.modules


def test_project_selection_edit_and_reopen_preserve_source(window, tmp_path):
    project, folder = make_project(window, tmp_path)
    original = (folder / "source" / "presentation.pptx").read_bytes()
    assert window.current_slide == 0
    window.page_text.setPlainText("修改后的第一页讲稿。")
    window.page_speed_follow.setChecked(False)
    window.page_speed.setValue(1.2)
    window.slide_list.setCurrentRow(1)
    assert project["slides"][0]["text"] == "修改后的第一页讲稿。"
    assert project["slides"][0]["speed"] == 1.2
    window.page_text.clear()
    window.page_blank.setValue(4)
    window.slide_list.item(2).setCheckState(Qt.CheckState.Unchecked)
    assert window.save_project()
    restored = projects.load_project(folder)
    assert restored["slides"][1]["text"] == ""
    assert restored["slides"][1]["blank_seconds"] == 4
    assert not restored["slides"][2]["included"]
    assert (folder / "source" / "presentation.pptx").read_bytes() == original
    window._set_project(restored, folder)
    assert window.page_text.toPlainText() == "修改后的第一页讲稿。"
    assert not window.test_errors


def test_text_generation_submits_real_reference_and_result_enables_exports(window, tmp_path, monkeypatch):
    requests = []
    monkeypatch.setattr(window.worker, "submit", lambda command, **payload: requests.append((command, payload)))
    window.text_edit.setPlainText("日期为 2026 年 9 月 9 日，欢迎使用。")
    window.generate_text()
    assert requests[0][0] == "synthesize"
    assert Path(requests[0][1]["ref_audio"]).is_file()
    assert requests[0][1]["ref_text"] == window.voice_items[0]["transcript"]
    assert requests[0][1]["seed"] == 0
    output = tmp_path / "completed.wav"
    shutil.copyfile(window.voice_items[0]["audio"], output)
    window._worker_event({"event": "result", "path": str(output), "duration": 12.78, "segments": 1})
    assert window.text_play.isEnabled()
    assert window.text_wav.isEnabled()
    assert window.text_mp3.isEnabled()
    assert window._text_output == str(output)


def test_cancelled_page_keeps_old_completed_audio(window, tmp_path, monkeypatch):
    project, folder = make_project(window, tmp_path)
    old = tmp_path / "old.wav"
    shutil.copyfile(window.voice_items[0]["audio"], old)
    projects.remember_audio(project, project["slides"][0], folder, old, "oldfingerprint")
    monkeypatch.setattr(window.worker, "submit", lambda command, **payload: "request")
    window._start_batch([0], force=True)
    assert window._active_job["kind"] == "page"
    window._worker_event({"event": "cancelled", "message": "任务已取消"})
    assert Path(projects.resolve_resource(folder, project["slides"][0]["audio"])).read_bytes() == old.read_bytes()
    assert project["slides"][0]["status"] == "pending"
    assert window._batch == []


def test_saved_temporary_voice_survives_original_file_removal(window, tmp_path):
    project, folder = make_project(window, tmp_path)
    source = tmp_path / "临时.wav"
    shutil.copyfile(window.voice_items[0]["audio"], source)
    window.ppt_voice.combo.setCurrentIndex(window.ppt_voice.combo.findData("__temporary__"))
    window.ppt_voice.audio.setText(str(source))
    window.ppt_voice.transcript.setPlainText(window.voice_items[0]["transcript"])
    assert window.save_project()
    source.unlink()
    restored = projects.load_project(folder)
    window._set_project(restored, folder)
    voice = window.ppt_voice.current()
    assert Path(voice["audio"]).is_file()
    assert not window.test_errors


def test_temporary_voice_identity_is_stable_between_saves(window, tmp_path):
    make_project(window, tmp_path)
    window.ppt_voice.combo.setCurrentIndex(window.ppt_voice.combo.findData("__temporary__"))
    window.ppt_voice.audio.setText(window.voice_items[0]["audio"])
    window.ppt_voice.transcript.setPlainText(window.voice_items[0]["transcript"])
    first = window.ppt_voice.current()["id"]
    assert window.ppt_voice.current()["id"] == first
    assert window.save_project()
    assert window.save_project()
    assert sum(voice_id == first for voice_id in window.project["voices"]) == 1
    window.ppt_voice.transcript.insertPlainText("新的原文")
    assert window.ppt_voice.current()["id"] != first


def test_text_draft_and_result_restore(window, tmp_path):
    window.text_edit.setPlainText("重启后仍然保留的讲稿。")
    window.text_voice.combo.setCurrentIndex(window.text_voice.combo.findData("__temporary__"))
    original = tmp_path / "临时录音.wav"
    shutil.copyfile(window.voice_items[0]["audio"], original)
    window.text_voice.audio.setText(str(original))
    window.text_voice.transcript.setPlainText(window.voice_items[0]["transcript"])
    output = tmp_path / "成品.wav"
    shutil.copyfile(original, output)
    window._text_output = str(output)
    window._save_text_state()
    original.unlink()
    window.text_edit.clear()
    window._text_output = None
    window._restore_text_state()
    assert window.text_edit.toPlainText() == "重启后仍然保留的讲稿。"
    assert Path(window.text_voice.current()["audio"]).is_file()
    assert window._text_output == str(output)


def test_unknown_gpu_cannot_bypass_gate_with_manual_paths(window, monkeypatch):
    called = []
    assert not MainWindow._ensure_inference_ready(window, lambda: called.append(True))
    assert not called
    assert window.test_errors


def test_signature_requires_real_selftest_after_driver_change(window, tmp_path, monkeypatch):
    import json
    from desktop_app.paths import user_data_dir
    runtime = tmp_path / "python.exe"
    runtime.write_bytes(b"runtime")
    model = tmp_path / "model"
    model.mkdir()
    (model / "cosyvoice3.yaml").write_text("test: model", encoding="utf-8")
    window.settings.update(runtime_python=str(runtime), model_dir=str(model), runtime_id="cu121")
    gpu = dict(uuid="GPU-TEST", driver="595.79", compute_capability="7.5", total_mb=8192, status="compatible")
    window.gpu_combo.blockSignals(True)
    window.gpu_combo.addItem("RTX 2070 Super", gpu)
    window.gpu_combo.setCurrentIndex(0)
    window.gpu_combo.blockSignals(False)
    calls = []
    monkeypatch.setattr(window, "self_test", lambda: calls.append("self_test"))
    assert not MainWindow._ensure_inference_ready(window, lambda: None)
    assert calls == ["self_test"]
    report = user_data_dir() / "last-self-test.json"
    report.write_text(json.dumps({"signature": window._environment_signature(), "inference_success": True}), encoding="utf-8")
    assert MainWindow._ensure_inference_ready(window, lambda: None)
    gpu["driver"] = "596.00"
    window.gpu_combo.setItemData(0, gpu)
    assert not MainWindow._ensure_inference_ready(window, lambda: None)
    assert calls == ["self_test", "self_test"]


def test_blank_pages_finish_without_starting_gpu(window, tmp_path, monkeypatch):
    project, folder = make_project(window, tmp_path)
    for slide in project["slides"]:
        slide["text"] = ""
    window._refresh_slides(keep=True)
    monkeypatch.setattr(window, "_ensure_inference_ready", lambda callback: pytest.fail("blank pages must not start a model"))
    window._start_batch([0, 1, 2])
    assert not window._batch
    assert not window.worker.busy
    assert not window.test_errors


def test_repair_chain_checks_all_managed_components_then_selftests(window, tmp_path, monkeypatch):
    from desktop_app import runtime
    gpu = dict(uuid="GPU-TEST", runtime_id="cu121", status="compatible", name="RTX 2070 SUPER", total_mb=8192)
    window.gpu_combo.blockSignals(True)
    window.gpu_combo.addItem("Test GPU", gpu)
    window.gpu_combo.setCurrentIndex(0)
    window.gpu_combo.blockSignals(False)
    monkeypatch.setattr(window, "apply_settings", lambda **kwargs: True)
    monkeypatch.setattr(window, "_refresh_components", lambda: {})
    calls = []
    monkeypatch.setattr(window, "self_test", lambda: calls.append("self_test"))
    def install(identifier, settings, progress, cancel):
        calls.append(identifier)
        return {"id": identifier, "path": str(tmp_path / identifier)}
    monkeypatch.setattr(runtime, "install_component", install)
    monkeypatch.setattr(window, "background", lambda message, function, success=None, failure=None: success(function(lambda event: None, lambda: False)))
    window.repair_components()
    assert calls == ["runtime-cu121", "model-cosyvoice3", "ffmpeg", "self_test"]


def test_generation_refreshes_stale_driver_probe_before_using_signature(window, monkeypatch):
    window._last_gpu_probe = time.monotonic() - 10
    calls = []
    monkeypatch.setattr(window, "detect_hardware", lambda: calls.append("probe"))
    callback = lambda: None
    assert not MainWindow._ensure_inference_ready(window, callback)
    assert calls == ["probe"]
    assert window._pending_action is callback
