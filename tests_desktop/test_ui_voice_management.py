"""Voice deletion, context recovery, empty-library and recording integration."""
import json
import os
from pathlib import Path
import shutil
import wave

if os.name == "nt":
    import ctypes
    ctypes.windll.kernel32.SetErrorMode(3)
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from PySide6.QtCore import QCoreApplication, QEvent, Qt
from PySide6.QtWidgets import QApplication, QDialog, QDialogButtonBox, QListWidget, QMessageBox
from pptx import Presentation

from desktop_app import paths, projects, ui
from tests_desktop.test_path_aliases import alias_directory


def seed_test_voice_library(data_root):
    voices = data_root / "voices" / "local-presets"
    voices.mkdir(parents=True, exist_ok=True)
    audio = voices / "reference.wav"
    with wave.open(str(audio), "wb") as stream:
        stream.setnchannels(1)
        stream.setsampwidth(2)
        stream.setframerate(24000)
        stream.writeframes((b"\x00\x10" * 24000))
    rows = [dict(id=f"test-local-{index:02d}", name=f"测试音色 {index:02d}",
                 reference_audio="reference.wav", transcript="测试参考原文",
                 distribution="local_only") for index in range(1, 7)]
    (voices / "manifest.json").write_text(
        json.dumps(dict(schema_version=1, voices=rows), ensure_ascii=False), encoding="utf-8"
    )


@pytest.fixture
def window(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])
    monkeypatch.setenv("COSYVOICE_DESKTOP_DATA", str(tmp_path / "data"))
    seed_test_voice_library(tmp_path / "data")
    session = tmp_path / "session"
    session.mkdir()
    monkeypatch.setattr(paths, "_SESSION", session)
    monkeypatch.setattr(ui.MainWindow, "detect_hardware", lambda self: None)
    result = ui.MainWindow()
    result._startup_timer.stop()
    errors = []
    monkeypatch.setattr(result, "show_error", errors.append)
    result.test_errors = errors
    monkeypatch.setattr(QMessageBox, "question", lambda *args, **kwargs: QMessageBox.StandardButton.Yes)
    yield result
    result.worker.active = None
    result._active_job = None
    result._background = None
    result._recording_dialog = None
    result._dirty = False
    result.close()
    result.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
    app.processEvents()


def presentation(tmp_path, text="本页讲稿"):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.notes_slide.notes_text_frame.text = text
    source = tmp_path / "中文测试.pptx"
    deck.save(source)
    return source


def hide_all(window):
    for voice in window.library.presets():
        window.library.delete(voice["id"])
    window._refresh_voice_list()


def test_delete_selected_preset_keeps_draft_project_selection_and_completed_audio(window, tmp_path):
    voice = window.voice_items[0]
    project = projects.create_project(presentation(tmp_path), tmp_path / "project", voice)
    projects.remember_audio(project, project["slides"][0], tmp_path / "project", voice["audio"], "stable")
    window._set_project(project, tmp_path / "project")
    window.text_voice.select(voice["id"])
    before = dict(project["slides"][0])
    window._refresh_voice_list(voice["id"])
    window.delete_voice()
    assert not window.test_errors
    assert voice["id"] not in {row["id"] for row in window.voice_items}
    assert Path(voice["audio"]).is_file()
    assert window.text_voice.current()["id"] == voice["id"]
    assert window.ppt_voice.current()["id"] == voice["id"]
    assert "草稿保留" in window.text_voice.combo.currentText()
    assert "工程保留" in window.ppt_voice.combo.currentText()
    assert project["slides"][0] == before
    assert not window._dirty
    assert window.voice_list.count() == len(window.library.list())


def test_delete_custom_copies_unsaved_references_before_removing_library_file(window, tmp_path):
    sample = window.voice_items[0]
    voice = window.library.add("将要删除", sample["audio"], sample["transcript"])
    window._refresh_voice_list(voice["id"])
    project = projects.create_project(presentation(tmp_path), tmp_path / "project")
    window._set_project(project, tmp_path / "project")
    window.text_voice.select(voice["id"])
    window.ppt_voice.select(voice["id"])
    window._refresh_voice_list(voice["id"])
    window.delete_voice()
    assert not window.test_errors
    assert not Path(voice["audio"]).exists()
    assert Path(window.text_voice.current()["audio"]).is_file()
    assert Path(window.ppt_voice.current()["audio"]).is_file()
    assert window.ppt_voice.current()["id"] == voice["id"]
    window._text_voice_snapshot = None
    window._restore_text_state()
    assert window.text_voice.current()["id"] == voice["id"]
    assert Path(window.text_voice.current()["audio"]).is_file()


def test_refresh_preserves_explicit_page_voice_without_marking_audio_stale(window, tmp_path):
    global_voice, page_voice = window.voice_items[:2]
    directory = tmp_path / "project"
    project = projects.create_project(presentation(tmp_path), directory, global_voice)
    project["slides"][0]["voice_id"] = page_voice["id"]
    projects.save_project(project, directory, {page_voice["id"]: page_voice})
    projects.remember_audio(project, project["slides"][0], directory, global_voice["audio"], "stable")
    window._set_project(project, directory)
    window._refresh_voice_list(page_voice["id"])
    window.delete_voice()
    assert window.page_voice.currentData() == page_voice["id"]
    assert project["slides"][0]["status"] == "complete"
    assert project["slides"][0]["fingerprint"] == "stable"


def test_ppt_temporary_reference_survives_delete_before_project_is_created(window):
    sample = window.voice_items[0]
    custom = window.library.add("临时引用的自建音色", sample["audio"], sample["transcript"])
    window._refresh_voice_list(custom["id"])
    window.ppt_voice.select("__temporary__")
    window.ppt_voice.audio.setText(custom["audio"])
    window.ppt_voice.transcript.setPlainText(custom["transcript"])
    original_id = window.ppt_voice.current()["id"]
    window.delete_voice()
    current = window.ppt_voice.current()
    assert current["id"] == original_id
    assert Path(current["audio"]).is_file()
    assert not Path(custom["audio"]).exists()
    assert not window.test_errors


def test_empty_library_clears_details_and_allows_blank_project_without_gpu(window, tmp_path, monkeypatch):
    hide_all(window)
    window.text_voice.select("")
    window.ppt_voice.select("")
    assert not window.voice_items
    assert window.voice_name.text() == "音色库为空"
    assert not window.voice_text.toPlainText()
    assert not window.voice_delete_button.isEnabled()
    assert not window.voice_preview_button.isEnabled()
    assert window.voice_record_button.isEnabled()
    assert window.voice_restore_button.isEnabled()
    project = projects.create_project(presentation(tmp_path, ""), tmp_path / "blank")
    window._set_project(project, tmp_path / "blank")
    assert window.save_project()
    monkeypatch.setattr(window, "_ensure_inference_ready", lambda *args: pytest.fail("blank pages need no inference"))
    window._start_batch([0])
    assert not window.test_errors
    assert not window.worker.busy
    assert project["settings"]["voice_id"] == ""


def test_empty_library_new_project_ui_imports_without_selecting_voice(window, tmp_path, monkeypatch):
    hide_all(window)
    window.ppt_voice.select("")
    source = presentation(tmp_path)
    monkeypatch.setattr(ui.QFileDialog, "getOpenFileName", lambda *args: (str(source), ""))
    monkeypatch.setattr(ui.QFileDialog, "getExistingDirectory", lambda *args: str(tmp_path))
    monkeypatch.setattr(ui.documents, "powerpoint_available", lambda: False)
    monkeypatch.setattr(window, "background", lambda message, function, callback: callback(function(lambda event: None, lambda: False)))
    window.new_project()
    assert window.project is not None
    assert window.project["settings"]["voice_id"] == ""
    assert (window.project_dir / "project.json").is_file()
    assert all("PowerPoint" in error for error in window.test_errors)


def test_missing_voice_blocks_synthesis_but_not_project_save(window, tmp_path, monkeypatch):
    hide_all(window)
    project = projects.create_project(presentation(tmp_path), tmp_path / "project")
    window._set_project(project, tmp_path / "project")
    assert window.save_project()
    monkeypatch.setattr(window.worker, "submit", lambda *args, **kwargs: pytest.fail("must require a voice first"))
    window._start_batch([0])
    assert "先选择音色或直接录音" in window.test_errors[-1]


def test_empty_library_self_test_uses_independent_worker_reference(window, monkeypatch):
    hide_all(window)
    monkeypatch.setattr(window, "apply_settings", lambda **kwargs: True)
    monkeypatch.setattr(window, "_environment_signature", lambda: "signature")
    window.gpu_combo.blockSignals(True)
    window.gpu_combo.addItem("test", {"status": "compatible"})
    window.gpu_combo.setCurrentIndex(window.gpu_combo.count() - 1)
    calls = []
    monkeypatch.setattr(window.worker, "submit", lambda command, **kwargs: calls.append((command, kwargs)))
    window.self_test()
    assert not window.test_errors
    assert calls[0][0] == "self_test"
    assert "ref_audio" not in calls[0][1]
    assert "ref_text" not in calls[0][1]


def test_restore_dialog_restores_only_checked_presets_without_switching_context(window, monkeypatch):
    target, other = window.voice_items[:2]
    window._preserve_voice_contexts()
    for voice in (target, other):
        window.library.delete(voice["id"])
    window._refresh_voice_list()
    before = window.text_voice.current()["id"]
    def accept_checked(dialog):
        choices = dialog.findChild(QListWidget)
        for index in range(choices.count()):
            item = choices.item(index)
            if item.data(Qt.ItemDataRole.UserRole) == other["id"]:
                item.setCheckState(Qt.CheckState.Checked)
        dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.StandardButton.Ok).click()
        return dialog.result()
    monkeypatch.setattr(ui.QDialog, "exec", accept_checked)
    window.restore_voices()
    visible = {voice["id"] for voice in window.voice_items}
    assert other["id"] in visible and target["id"] not in visible
    assert window.text_voice.current()["id"] == before


def test_legacy_hidden_named_draft_is_migrated_without_changing_voice(window):
    voice = window.voice_items[1]
    directory = paths.user_data_dir() / "text-session"
    directory.mkdir()
    (directory / "draft.json").write_text(json.dumps({"text": "旧讲稿", "voice_id": voice["id"], "speed": 1.2}), encoding="utf-8")
    window.library.delete(voice["id"])
    window._refresh_voice_list()
    window._restore_text_state()
    assert window.text_voice.current()["id"] == voice["id"]
    assert window.text_edit.toPlainText() == "旧讲稿"
    saved = json.loads((directory / "draft.json").read_text(encoding="utf-8"))
    assert saved["schema_version"] == 2
    assert not Path(saved["voice"]["audio"]).is_absolute()
    assert (directory / saved["voice"]["audio"]).is_file()


def test_legacy_missing_voice_never_falls_back_to_another_speaker(window):
    directory = paths.user_data_dir() / "text-session"
    directory.mkdir()
    (directory / "draft.json").write_text(json.dumps({"text": "保留文字", "voice_id": "custom-gone"}), encoding="utf-8")
    window._restore_text_state()
    assert window.text_voice.combo.currentData() == "custom-gone"
    assert "不可用" in window.text_voice.combo.currentText()
    assert window.text_voice.current(required=False) is None
    assert window.text_edit.toPlainText() == "保留文字"


def test_context_copy_failure_prevents_delete(window, monkeypatch):
    voice = window.voice_items[0]
    with monkeypatch.context() as patch:
        patch.setattr(window, "_save_text_state", lambda **kwargs: (_ for _ in ()).throw(OSError("磁盘写入失败")))
        window.delete_voice()
    assert voice["id"] in {item["id"] for item in window.library.list()}
    assert "磁盘写入失败" in window.test_errors[-1]


def test_busy_methods_cannot_start_recording_delete_or_restore(window):
    before = [voice["id"] for voice in window.voice_items]
    window.worker.active = {"id": "running"}
    window.update_busy()
    assert not window.voice_record_button.isEnabled()
    assert not window.text_voice.record_button.isEnabled()
    assert not window.ppt_voice.record_button.isEnabled()
    window.delete_voice()
    window.restore_voices()
    window.record_voice()
    window.record_temporary_voice(window.text_voice)
    assert len(window.test_errors) == 4
    assert [voice["id"] for voice in window.library.list()] == before


@pytest.mark.parametrize("temporary", [False, True])
def test_recording_result_is_copied_before_dialog_cleanup(window, tmp_path, monkeypatch, temporary):
    from desktop_app import recording
    sample = window.voice_items[0]
    original = tmp_path / "dialog.wav"
    shutil.copyfile(sample["audio"], original)
    calls = []
    class Dialog:
        def __init__(self, parent, mode, stop_playback):
            calls.append(mode)
            stop_playback()
        def exec(self):
            return QDialog.DialogCode.Accepted
        def take_result(self):
            return {"audio": str(original), "transcript": sample["transcript"], "name": "录制测试"}
        def cleanup(self):
            original.unlink(missing_ok=True)
            calls.append("cleanup")
        def deleteLater(self):
            calls.append("deleted")
    monkeypatch.setattr(recording, "RecordingDialog", Dialog)
    if temporary:
        window.record_temporary_voice(window.text_voice)
        voice = window.text_voice.current()
        assert voice["kind"] == "temporary"
        assert voice["audio"] != str(original)
        assert len(window.library.list()) == 6
    else:
        window.record_voice()
        voice = next(voice for voice in window.library.list() if voice["name"] == "录制测试")
        assert voice["source"] == "软件录制"
    assert Path(voice["audio"]).is_file()
    assert not original.exists()
    assert calls == ["temporary" if temporary else "library", "cleanup", "deleted"]
    assert window._recording_dialog is None
    assert not window.test_errors
    assert not (paths.user_data_dir() / "recording-recovery").exists()


@pytest.mark.parametrize("temporary", [False, True])
def test_failed_recording_save_preserves_audio_and_text_without_changing_voice(window, tmp_path, monkeypatch, temporary):
    from desktop_app import recording
    sample = window.text_voice.current()
    original = tmp_path / "accepted-recording.wav"
    shutil.copyfile(sample["audio"], original)
    expected_hash = ui.documents.sha256(original)
    calls = []
    class Dialog:
        def __init__(self, *args, **kwargs):
            pass
        def exec(self):
            return QDialog.DialogCode.Accepted
        def take_result(self):
            return {"audio": str(original), "transcript": sample["transcript"], "name": "等待恢复的录音"}
        def cleanup(self):
            original.unlink(missing_ok=True)
            calls.append("cleanup")
        def deleteLater(self):
            calls.append("deleted")
    def fail(*args, **kwargs):
        raise OSError("无法保存音色")
    monkeypatch.setattr(recording, "RecordingDialog", Dialog)
    with monkeypatch.context() as patch:
        if temporary:
            patch.setattr(window, "_save_text_state", fail)
            window.record_temporary_voice(window.text_voice)
        else:
            patch.setattr(window.library, "add", fail)
            window.record_voice()
    recovery_dirs = list((paths.user_data_dir() / "recording-recovery").iterdir())
    assert len(recovery_dirs) == 1
    recovery = recovery_dirs[0]
    saved = json.loads((recovery / "recording.json").read_text(encoding="utf-8"))
    assert saved["name"] == "等待恢复的录音"
    assert saved["transcript"] == sample["transcript"]
    assert ui.documents.sha256(recovery / saved["audio"]) == expected_hash
    assert str(recovery) in window.test_errors[-1]
    assert window.text_voice.current()["id"] == sample["id"]
    assert len(window.library.list()) == 6
    assert not original.exists()
    assert calls == ["cleanup", "deleted"]
    assert window._recording_dialog is None


def test_recovery_storage_failure_does_not_delete_only_recording(window, tmp_path, monkeypatch):
    from desktop_app import recording
    sample = window.voice_items[0]
    original = tmp_path / "only-recording.wav"
    shutil.copyfile(sample["audio"], original)
    calls = []
    class Dialog:
        def __init__(self, *args, **kwargs):
            pass
        def exec(self):
            return QDialog.DialogCode.Accepted
        def take_result(self):
            return {"audio": str(original), "transcript": sample["transcript"], "name": "只此一份"}
        def cleanup(self):
            original.unlink()
            calls.append("cleanup")
        def deleteLater(self):
            calls.append("deleted")
    def fail(*args, **kwargs):
        raise OSError("磁盘不可写")
    monkeypatch.setattr(recording, "RecordingDialog", Dialog)
    with monkeypatch.context() as patch:
        patch.setattr(window.library, "add", fail)
        patch.setattr(ui.shutil, "copyfile", fail)
        window.record_voice()
    assert original.is_file()
    assert str(original) in window.test_errors[-1]
    assert sample["transcript"] in window.test_errors[-1]
    assert "关闭软件前复制" in window.test_errors[-1]
    assert calls == ["deleted"]
    assert window._recording_dialog is None


@pytest.mark.parametrize("temporary", [False, True])
def test_text_draft_snapshots_round_trip_with_virtualized_appdata_directory(window, alias_directory, monkeypatch, temporary):
    alias, physical = alias_directory
    # Windows AppData virtualization also redirects creation through a symlink;
    # create this existing-profile directory at its physical location first.
    (physical / "text-session").mkdir()
    monkeypatch.setattr(ui, "user_data_dir", lambda: alias)
    sample = window.voice_items[0]
    if temporary:
        window.text_voice.select("__temporary__")
        window.text_voice.audio.setText(sample["audio"])
        window.text_voice.transcript.setPlainText(sample["transcript"])
    expected_id = window.text_voice.current()["id"]
    window.text_edit.setPlainText("重启后保留这份讲稿与同一个声音。")
    assert window._save_text_state(strict=True)
    saved = json.loads((physical / "text-session" / "draft.json").read_text(encoding="utf-8"))
    assert not Path(saved["voice"]["audio"]).is_absolute()
    if temporary:
        assert saved["temporary"]["audio"] == saved["voice"]["audio"]
    window._text_voice_snapshot = None
    window._restore_text_state()
    restored = window.text_voice.current()
    assert restored["id"] == expected_id
    assert Path(restored["audio"]).is_file()
    assert "重启后保留" in window.text_edit.toPlainText()
    assert not window.test_errors
